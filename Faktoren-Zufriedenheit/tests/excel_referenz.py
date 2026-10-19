#!/usr/bin/env python3
"""Referenzwerte aus dem Excel: Zufallszustaende, Neuberechnung mit LibreOffice und
unabhaengige Nachrechnung in Python (Decimal, ohne Gleitkomma).

Wird von pruefe_excel.py und test_ui_excel.py benutzt. Nicht direkt ausfuehren.

Vorgehen
  1. Das Original-Excel (quellen/p2/gluecksfaktoren.xlsx) wird geladen, Formeln bleiben unveraendert.
  2. In die Eingabezellen G-M (Werte vor der Entscheidung) und O (Wuerfel) der Zeilen 3-46 werden
     Zufallswerte geschrieben. Je Durchgang bekommt jede der 38 Zeilen einen eigenen Zustand.
  3. LibreOffice berechnet die Datei neu (headless). Gelesen werden Q, S-Y, AA-AF, AI, AN.
  4. Variante «original»: unveraendert (Y3 und Y40 sind im Quell-Excel leer).
     Variante «repariert»: Y3 und Y40 bekommen dieselbe Formel wie die anderen 36 Zeilen.
  5. Zusaetzlich rechnet rechne_decimal() jede Zeile unabhaengig mit exakter Dezimalarithmetik nach.
"""
import json
import os
import random
import shutil
import subprocess
import tempfile
from decimal import Decimal, ROUND_HALF_UP, getcontext

import openpyxl

QUELLE_XLSX = "/home/claude/html-projekte/quellen/p2/gluecksfaktoren.xlsx"
ZEILEN = list(range(3, 11)) + list(range(13, 21)) + list(range(23, 31)) + list(range(33, 47))
EINGABE = [("zeit", "G"), ("energie", "H"), ("soz", "I"), ("erfuellung", "J"),
           ("lebensstil", "K"), ("finanzen", "L"), ("stimmung", "M")]
AUSGABE = {  # Schluessel -> Spalte
    "resultat": "Q",
    "nachher_zeit": "S", "nachher_energie": "T", "nachher_soz": "U", "nachher_erfuellung": "V",
    "nachher_lebensstil": "W", "nachher_finanzen": "X", "nachher_stimmung": "Y",
    "wirk_zeit": "AA", "wirk_energie": "AB", "wirk_soz": "AC", "wirk_erfuellung": "AD",
    "wirk_lebensstil": "AE", "wirk_finanzen": "AF",
    "energie_aenderung": "AI", "stimmung_aenderung": "AN",
}
Y_FORMEL = '=IF(O{r}<>"", M{r}+AN{r}, "")'


# ------------------------------------------------------------------ Zufallszustaende
def halbschritte(rng, kante):
    """Zahl von -10 bis +10 in Schritten von 0.5; bevorzugt Grenzen und die Schwellen 2 / 4 / -2."""
    kanten = [-10, -9.5, -5, -3, -2.5, -2, -1.5, -0.5, 0, 0.5, 1, 1.5, 2, 2.5, 3, 4, 4.5, 5, 9.5, 10]
    if rng.random() < kante:
        return rng.choice(kanten)
    return rng.randint(-20, 20) / 2


def zufallszustand(rng):
    s = {k: halbschritte(rng, 0.35) for k, _ in EINGABE[:6]}
    r = rng.random()
    if r < 0.15:
        s["stimmung"] = rng.choice([0, 1, 2, 3, 99, 100])
    elif r < 0.25:
        s["stimmung"] = rng.randint(-8, -1)
    else:
        s["stimmung"] = rng.randint(0, 100)
    s["wurf"] = rng.choice([1, 20, rng.randint(1, 20), rng.randint(1, 20), rng.randint(1, 20)])
    return s


def zahl_oder_ganz(x):
    return int(x) if float(x) == int(x) else float(x)


def erzeuge_faelle(durchgaenge, seed=20261013):
    """Liste von Faellen: {'durchgang': d, 'zeile': r, 'zustand': {...}}"""
    rng = random.Random(seed)
    faelle = []
    for d in range(durchgaenge):
        for r in ZEILEN:
            z = zufallszustand(rng)
            faelle.append({"durchgang": d, "zeile": r, "zustand": {k: zahl_oder_ganz(v) for k, v in z.items()}})
    return faelle


# ------------------------------------------------------------------ LibreOffice
def libreoffice_rechnen(xlsx_in, ausgabeordner):
    profil = tempfile.mkdtemp(prefix="lo-profil-")
    try:
        cmd = ["soffice", "-env:UserInstallation=file://" + profil, "--headless", "--norestore",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", ausgabeordner, xlsx_in]
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        if p.returncode != 0:
            raise RuntimeError("LibreOffice-Fehler: " + p.stdout + p.stderr)
    finally:
        shutil.rmtree(profil, ignore_errors=True)
    return os.path.join(ausgabeordner, os.path.basename(xlsx_in))


def excel_neu_rechnen(faelle, variante, arbeitsordner):
    """Schreibt je Durchgang eine Kopie des Excels mit den Eingaben, rechnet mit LibreOffice neu und
    liest die Ausgaben. Gibt dict {(durchgang, zeile): {...}} zurueck."""
    durchgaenge = sorted({f["durchgang"] for f in faelle})
    ergebnis = {}
    for d in durchgaenge:
        wb = openpyxl.load_workbook(QUELLE_XLSX)
        ws = wb["Tabelle1"]
        if variante in ("repariert", "sequentiell"):
            for r in (3, 40):
                assert ws["Y%d" % r].value is None
                ws["Y%d" % r].value = Y_FORMEL.format(r=r)
        elif variante != "original":
            raise ValueError(variante)
        if variante == "sequentiell":
            # Nur die SUMME in Q wird durch eine ausgeschriebene Addition ersetzt (AA+AB+...+AF).
            # Sonst identische Formel. Grund: LibreOffice summiert SUM() in anderer Reihenfolge als
            # eine einfache Schleife; bei exakten Haelften kann das Rauschen die Rundung kippen.
            for r in ZEILEN:
                alt = ws["Q%d" % r].value
                neu = alt.replace("SUM(AA%d:AF%d)" % (r, r), "(AA%d+AB%d+AC%d+AD%d+AE%d+AF%d)" % ((r,) * 6))
                assert neu != alt
                ws["Q%d" % r].value = neu
        # Eingaben in alle 38 Zeilen
        for f in faelle:
            if f["durchgang"] != d:
                continue
            r = f["zeile"]
            for k, sp in EINGABE:
                ws["%s%d" % (sp, r)].value = f["zustand"][k]
            ws["O%d" % r].value = f["zustand"]["wurf"]
        # Zeile 4 teilt sich die Startwerte mit Zeile 3 im Quell-Excel; hier ueberschrieben, das ist gewollt.
        ein = os.path.join(arbeitsordner, "%s_%02d.xlsx" % (variante, d))
        wb.save(ein)
        aus = os.path.join(arbeitsordner, "rechnung_" + variante)
        os.makedirs(aus, exist_ok=True)
        fertig = libreoffice_rechnen(ein, aus)
        wr = openpyxl.load_workbook(fertig, data_only=True)["Tabelle1"]
        for f in faelle:
            if f["durchgang"] != d:
                continue
            r = f["zeile"]
            werte = {}
            for k, sp in AUSGABE.items():
                v = wr["%s%d" % (sp, r)].value
                werte[k] = v
            ergebnis[(d, r)] = werte
    return ergebnis


# ------------------------------------------------------------------ unabhaengige Nachrechnung
getcontext().prec = 50


def D(x):
    return Decimal(str(x))


def runde(x, stellen):
    """Excel-ROUND mit exakter Dezimalarithmetik: halb weg von null."""
    q = Decimal(1).scaleb(-stellen)
    if x >= 0:
        return x.quantize(q, rounding=ROUND_HALF_UP)
    return -((-x).quantize(q, rounding=ROUND_HALF_UP))


def lies_zeilenparameter():
    """Zeilenparameter direkt aus dem Excel (Konstanten), als Decimal."""
    ws = openpyxl.load_workbook(QUELLE_XLSX)["Tabelle1"]
    out = {}
    for r in ZEILEN:
        def c(sp):
            return D(ws["%s%d" % (sp, r)].value)
        out[r] = {
            "P": c("P"), "AH": c("AH"), "AJ": c("AJ"), "AK": c("AK"), "AL": c("AL"), "AM": c("AM"),
            "AP": D(round(ws["AP%d" % r].value, 10)), "AQ": D(round(ws["AQ%d" % r].value, 10)),
            "AR": D(round(ws["AR%d" % r].value, 10)), "AS": D(round(ws["AS%d" % r].value, 10)),
            "AT": D(round(ws["AT%d" % r].value, 10)), "AU": D(round(ws["AU%d" % r].value, 10)),
            "AW": c("AW"),
        }
    return out


def rechne_decimal(p, z):
    """Die Excel-Formeln, Zelle fuer Zelle, in exakter Dezimalarithmetik."""
    G, H, I, J, K, L, M = (D(z[k]) for k in ("zeit", "energie", "soz", "erfuellung", "lebensstil", "finanzen", "stimmung"))
    O = D(z["wurf"])

    def sat(wert, w):  # =IF(G>2, w*2, IF(G>4, w*3, G*w))
        if wert > 2:
            return w * 2
        if wert > 4:
            return w * 3
        return wert * w

    AA = sat(G, p["AP"])
    AB = sat(H, p["AQ"])
    AC = p["AR"] * I
    AD = p["AS"] * J
    AE = p["AT"] * K
    AF = sat(L, p["AU"])
    summe = AA + AB + AC + AD + AE + AF
    q_roh = runde(2 * summe + O, 1)
    Q = Decimal(0) if q_roh < 0 else q_roh
    AN = runde(8 * (Q - p["P"]) / 10, 0)
    Y = M + AN
    AI = p["AW"] + runde(8 * (Y - 50) / 50, 0) / 2
    S = G + p["AH"] - 1 if L < -2 else G + p["AH"]
    return {
        "resultat": Q, "stimmung_aenderung": AN, "nachher_stimmung": Y, "energie_aenderung": AI,
        "nachher_zeit": S, "nachher_energie": H + AI, "nachher_soz": I + p["AJ"],
        "nachher_erfuellung": J + p["AK"], "nachher_lebensstil": K + p["AL"], "nachher_finanzen": L + p["AM"],
        "wirk_zeit": AA, "wirk_energie": AB, "wirk_soz": AC, "wirk_erfuellung": AD,
        "wirk_lebensstil": AE, "wirk_finanzen": AF,
    }


def ist_haelftenfall(params, z):
    """True, wenn 2*SUMME+Wuerfel exakt auf x.x5 endet (Rundung auf eine Stelle entscheidet ueber die Haelfte)."""
    G, H, I, J, K, L = (D(z[k]) for k in ("zeit", "energie", "soz", "erfuellung", "lebensstil", "finanzen"))

    def sat(w, g):
        return g * 2 if w > 2 else w * g

    s = sat(G, params["AP"]) + sat(H, params["AQ"]) + params["AR"] * I + params["AS"] * J + params["AT"] * K + sat(L, params["AU"])
    v = 2 * s + D(z["wurf"])
    return (abs(v) * 100) % 10 == 5


def erzeuge_haelftenfaelle(durchgaenge, seed):
    rng = random.Random(seed)
    params = lies_zeilenparameter()
    faelle = []
    for d in range(durchgaenge):
        for r in ZEILEN:
            while True:
                z = zufallszustand(rng)
                if ist_haelftenfall(params[r], z):
                    break
            faelle.append({"durchgang": d, "zeile": r, "zustand": {k: zahl_oder_ganz(v) for k, v in z.items()}})
    return faelle


def referenz_erzeugen(durchgaenge=12, ordner=None, seed=20261013, nur_haelften=False, varianten=("original", "repariert", "sequentiell")):
    """Gibt (faelle, {variante: ergebnisse}, decimal) zurueck. Schluessel (durchgang, zeile)."""
    arbeits = ordner or tempfile.mkdtemp(prefix="p2-excel-")
    os.makedirs(arbeits, exist_ok=True)
    faelle = erzeuge_haelftenfaelle(durchgaenge, seed) if nur_haelften else erzeuge_faelle(durchgaenge, seed)
    lo = {v: excel_neu_rechnen(faelle, v, arbeits) for v in varianten}
    params = lies_zeilenparameter()
    dec = {}
    for f in faelle:
        dec[(f["durchgang"], f["zeile"])] = {k: float(v) for k, v in rechne_decimal(params[f["zeile"]], f["zustand"]).items()}
    return faelle, lo, dec


if __name__ == "__main__":
    raise SystemExit("Nicht direkt ausfuehren; siehe pruefe_excel.py")
