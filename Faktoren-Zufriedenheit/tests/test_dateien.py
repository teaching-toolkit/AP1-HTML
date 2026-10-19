"""Datenpruefung ohne Browser: Vollstaendigkeit, unabhaengiger Abgleich mit Word und Excel, Reproduzierbarkeit des Generators,
Dateigroesse, keine externen Adressen, kein Eszett.   Aufruf: python3 -I tests/test_dateien.py"""
import json
import os
import re
import subprocess
import sys
import tempfile
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx
import openpyxl

import hilfen as H
import referenz_spiel as R

HIER = os.path.dirname(os.path.abspath(__file__))
PROJ = H.PROJEKT
WORD = "/home/claude/html-projekte/quellen/p2/gluecksfaktor_spiel.docx"
E = H.Ergebnisse()
G = "-"


def lade_daten():
    t = open(os.path.join(PROJ, "spiel-daten.js"), encoding="utf-8").read()
    return json.loads(t[t.index("{"):t.rindex("}") + 1])


def lade_texte():
    """spiel-texte.js ist von Hand geschrieben, aber JSON-kompatibel (Schluessel in Anfuehrungszeichen): ohne Node lesbar."""
    t = open(os.path.join(PROJ, "spiel-texte.js"), encoding="utf-8").read()
    m = re.search(r"var SPIEL_TEXTE\s*=\s*", t)
    return json.loads(t[m.end():].rstrip().rstrip(";"))


def norm(s):
    s = re.sub(r"[„“”\"'«»]", "", s)
    s = re.sub(r"\s*/\s*", "/", s)
    return re.sub(r"\s+", " ", s.replace("\u2212", "-").replace("\u00a0", " ")).strip()


def word_lesen():
    paras = [(p.style.name, p.text.strip()) for p in docx.Document(WORD).paragraphs]
    eps, minis, cur, i = {}, {}, None, 0
    ATTR = {"Zeit": "zeit", "Energie": "energie", "Soziale Verbindungen": "soz", "Persönliche Erfüllung": "erfuellung", "Lebensstil-Qualität": "lebensstil", "Einkommen und Vermögen": "finanzen"}
    mode = None
    for st, tx in paras:
        m = re.match(r"^Episode (\w+)\s*[:\u2013-]\s*(.+)$", tx) if st == "Heading 2" else None
        if m:
            cur = eps.setdefault(m.group(1), {"titel": m.group(2), "opt": {}, "glueck": [], "erk": []}); mode = "ep"; opt = None; gl = False; continue
        m = re.match(r"^(M\d+): (.+)$", tx)
        if m and st.startswith("Normal"):
            cur = minis.setdefault(m.group(1), {"titel": m.group(2), "opt": {}, "glueck": ""}); mode = "mikro"; opt = None; continue
        if st.startswith("Heading") and not re.match(r"^Episode", tx) and mode:
            if tx.startswith("(nur nach") or re.match(r"^BLOCK \d", tx, re.I):
                continue
            mode = None; cur = None; continue
        if cur is None or not tx:
            if not tx:
                gl = False
            continue
        if mode == "ep":
            if tx.startswith("Vorlesen:"):
                rest = tx.split(":", 1)[1].strip()
                if rest:
                    cur["vorlesen"] = rest
                cur["_v"] = True; continue
            if cur.get("_v") and "vorlesen" not in cur:
                cur["vorlesen"] = tx; continue
            mo = re.match(r"^([AB]) (.+)$", tx)
            if mo and not tx.startswith("Würfeln"):
                opt = cur["opt"].setdefault(mo.group(1), {"label": mo.group(2)}); gl = False; continue
            if tx.startswith("Unmittelbare Folge:") and opt is not None:
                opt["folge"] = tx.split(":", 1)[1].strip(); continue
            if tx.startswith("Würfeln") and opt is not None:
                opt["wuerfel"] = tx[len("Würfeln"):].strip(); continue
            if tx.startswith("Weiter bei:") and opt is not None:
                mw = re.search(r"Episode (\w+)", tx)
                opt["weiter"] = mw.group(1) if mw else None; continue
            if tx.startswith("Glücksfaktoren"):
                gl = True; continue
            if gl:
                (cur["erk"] if re.match(r"^(Wichtige |Forschungs)[Ee]rkenntnis", tx) else cur["glueck"]).append(tx)
        elif mode == "mikro":
            if tx.startswith("Vorlesen:"):
                cur["vorlesen"] = tx.split(":", 1)[1].strip(); continue
            mo = re.match(r"^([AB]) (.+?) → (.+?)( → .*)?$", tx)
            if mo:
                d = {}
                for teil in re.split(r",\s*", mo.group(3).rstrip(". ")):
                    mt = re.match(r"^(.+?) ([+\-−]?\d+)\.?$", teil.strip())
                    d[ATTR[mt.group(1)]] = int(mt.group(2).replace("−", "-"))
                cur["opt"][mo.group(1)] = {"label": mo.group(2), "deltas": d}; continue
            if tx.startswith("Glücksfaktor"):
                cur["glueck"] = tx.split(":", 1)[1].strip()
    return eps, minis


WORTE = lambda s: {w for w in re.findall(r"[a-zäöü\u00df]+", s.lower()) if len(w) > 3}
BEHAUPTEND = [r"führt", r"steigert", r"verbessert", r"erhöht", r"wissenschaft", r"\bstudien?\b", r"bewiesen", r"senkt", r"macht glücklich", r"Glück", r"nachweislich", r"fördert"]


def texte_pruefen(d, eps_d, ids):
    """Handgeschriebene Datei spiel-texte.js gegen die erzeugten Daten: Zuordnung, Ausrichtung, Kandidaten (Ablenker)."""
    T = lade_texte()
    E.pruefe("spiel-texte.js: Episoden = Episoden der erzeugten Daten, je Episode Ausrichtung, Ablenker und Ergebnistexte", G,
             sorted(T["episoden"]) == sorted(ids) and all({"ausrichtung", "plausibel", "unplausibel", "ergebnis", "unscharf"} <= set(T["episoden"][i]) for i in ids), "%d Episoden" % len(T["episoden"]))
    # Ausrichtung: je Episode genau eine langfristig, eine kurzfristig; gegen die unabhaengige Tabelle
    fehl = []
    for i in ids:
        a = T["episoden"][i]["ausrichtung"]
        if sorted(a.values()) != ["kurzfristig", "langfristig"]: fehl.append(i + " nicht je eine")
        if a != R.AUSRICHTUNG[i]: fehl.append(i + " weicht von der Referenztabelle ab")
        if bool(T["episoden"][i]["unscharf"]) != (i in R.UNSCHARF): fehl.append(i + " unscharf")
        if T["episoden"][i]["unscharf"] and len(T["episoden"][i]["notiz"]) < 30: fehl.append(i + " Notiz fehlt")
    E.pruefe("Ausrichtung: je Episode eine langfristige und eine kurzfristige Option (19 x), gleich der Referenztabelle; unscharfe Episoden %s begruendet" % sorted(R.UNSCHARF), G, not fehl, "; ".join(fehl[:4]))
    E.pruefe("Ausrichtung: Ausnahme 11B (A langfristig, B kurzfristig) und sonst A kurzfristig, B langfristig", G,
             T["episoden"]["11B"]["ausrichtung"] == {"A": "langfristig", "B": "kurzfristig"} and all(T["episoden"][i]["ausrichtung"] == {"A": "kurzfristig", "B": "langfristig"} for i in ids if i != "11B"), "")
    # Kandidaten
    fehl, nstat, anzahlen = [], Counter(), []
    alle_echt = {}
    for i in ids:
        alle_echt[i] = [norm(re.sub(r"^[A-G]\d\s*[:\u2013-]\s*", "", re.sub(r"\s*\([^)]*\)\s*$", "", g))) for g in eps_d[i]["glueck"]]
    alle_echt_namen = {n for v in alle_echt.values() for n in v}
    gesehen_global = Counter()
    for i in ids:
        t = T["episoden"][i]
        echt = alle_echt[i]
        pl, un = t["plausibel"], t["unplausibel"]
        n = len(echt) + len(pl) + len(un)
        anzahlen.append(n)
        if not (6 <= n <= 8): fehl.append("%s: %d Eintraege" % (i, n))
        if len(echt) < 1 or len(pl) < 2 or len(un) < 2: fehl.append("%s: zu wenige echte/plausible/unplausible (%d/%d/%d)" % (i, len(echt), len(pl), len(un)))
        namen = [norm(x).lower() for x in echt + pl + un]
        if len(set(namen)) != len(namen): fehl.append(i + ": doppelter Eintrag")
        for x in pl + un:
            gesehen_global[x] += 1
            if len(x) > 70 or len(x) < 8: fehl.append("%s: Laenge %r" % (i, x))
            if chr(0xDF) in x: fehl.append(i + " Eszett")
            if any(re.search(b, x, re.I) for b in BEHAUPTEND): fehl.append("%s: behauptende Formulierung %r" % (i, x))
            # kein Ablenker gleicht oder ueberlappt stark mit einem echten Faktor dieser Episode oder irgendeiner Episode
            for e_ in alle_echt_namen:
                a, b = WORTE(x), WORTE(e_)
                if a and b and len(a & b) / len(a | b) >= 0.5: fehl.append("%s: %r aehnelt dem echten Faktor %r" % (i, x, e_))
        if not all(x.strip() == x and x[0].isupper() for x in pl + un): fehl.append(i + ": Schreibweise")
    E.pruefe("Ablenker: 6 bis 8 Eintraege je Episode, mindestens 2 plausible und 2 unplausible, keine Doppelten, keine Behauptungen, keiner aehnelt einem echten Faktor", G, not fehl,
             "; ".join(fehl[:4]) if fehl else "Anzahl je Episode %d bis %d; %d plausible, %d unplausible Ablenker" % (min(anzahlen), max(anzahlen),
                                                                                  sum(len(T['episoden'][i]['plausibel']) for i in ids), sum(len(T['episoden'][i]['unplausibel']) for i in ids)))
    wiederholt = [x for x, c in gesehen_global.items() if c > 1]
    E.pruefe("Ablenker: kein Ablenker kommt in zwei Episoden vor (kein Muster «immer wieder derselbe Ablenker»)", G, not wiederholt, str(wiederholt[:4]))
    # echte Faktoren: Kurznamen aus den Daten sinnvoll
    ok = all(2 < len(n) <= 70 and not re.match(r"^[A-G]\d", n) for v in alle_echt.values() for n in v)
    E.pruefe("Echte Faktoren: Kurzname (ohne Code und Klammerzusatz) aus den Daten, 3 bis 70 Zeichen", G, ok, "%d echte Faktoren in 19 Episoden" % sum(len(v) for v in alle_echt.values()))
    # Option-Buchstaben im Wort der Daten (werden bei Vertauschung umgeschrieben)
    mit_opt = [(i, g) for i in ids for g in eps_d[i]["glueck"] if re.search(r"Option [AB]", g)]
    E.pruefe("Hinweis: Glücksfaktoren mit Option-Buchstaben im Wort (werden bei vertauschten Buchstaben umgeschrieben)", G, True, "%d Eintraege, Episoden %s" % (len(mit_opt), sorted({i for i, _ in mit_opt})))


def main():
    d = lade_daten()
    eps_d = {e["id"]: e for e in d["episoden"]}
    mik_d = {m["id"]: m for m in d["mikroszenen"]}
    ids = ["1", "2A", "2B", "3", "4", "5A", "5B", "6", "7", "8A", "8B", "9", "10", "11A", "11B", "12A1", "12A2", "12B1", "12B2"]
    E.pruefe("19 Episoden in der verlangten Reihenfolge-Menge, 38 Optionen", G, sorted(eps_d) == sorted(ids) and sum(len(e["optionen"]) for e in d["episoden"]) == 38, "%d Episoden" % len(eps_d))
    E.pruefe("10 Mikroszenen M1-M10 mit je zwei Optionen", G, sorted(mik_d, key=lambda s: int(s[1:])) == ["M%d" % i for i in range(1, 11)] and all(set(m["optionen"]) == {"A", "B"} for m in d["mikroszenen"]), "")
    # Word unabhaengig
    eps, minis = word_lesen()
    E.pruefe("Word: 19 Episoden und 10 Mikroszenen unabhaengig gelesen", G, sorted(eps) == sorted(ids) and len(minis) == 10, "%d / %d" % (len(eps), len(minis)))
    fehl = []
    for i in ids:
        w, e = eps[i], eps_d[i]
        if norm(w["titel"]) != norm(e["titel"]): fehl.append(i + " Titel")
        if norm(w.get("vorlesen", "")) != norm(e["vorlesen"]): fehl.append(i + " Vorlesen")
        if [norm(x) for x in w["glueck"]] != [norm(x) for x in e["glueck"]]: fehl.append(i + " Glücksfaktoren")
        if [norm(x) for x in w["erk"]] != [norm(x) for x in e["erkenntnis"]]: fehl.append(i + " Erkenntnis")
        for k in "AB":
            wo, eo = w["opt"][k], e["optionen"][k]
            if norm(wo["label"]) != norm(eo["label"]): fehl.append("%s%s Label" % (i, k))
            if norm(wo.get("folge", "")) != norm(eo["folge"]): fehl.append("%s%s Folge" % (i, k))
            if norm(wo.get("wuerfel", "")) != norm(eo["wuerfeltitel"]): fehl.append("%s%s Wuerfeltitel" % (i, k))
            if wo.get("weiter") != eo["weiter"]: fehl.append("%s%s Weiter bei: Word %s, Daten %s" % (i, k, wo.get("weiter"), eo["weiter"]))
            if R.WEITER.get((i, k)) != eo["weiter"]: fehl.append("%s%s Weiter (Referenztabelle)" % (i, k))
    E.pruefe("Word gegen Daten: Titel, Vorlesen, 38 Optionen (Text, Folge, Wuerfeltitel), Weiter bei, Glücksfaktoren, Erkenntnis", G, not fehl, "; ".join(fehl[:5]))
    fehl = []
    for mid, w in minis.items():
        m = mik_d[mid]
        if norm(w["titel"]) != norm(m["titel"]): fehl.append(mid + " Titel")
        if norm(w.get("vorlesen", "")) != norm(m["vorlesen"]): fehl.append(mid + " Vorlesen")
        for k in "AB":
            if norm(w["opt"][k]["label"]) != norm(m["optionen"][k]["label"]): fehl.append(mid + k + " Label")
            if {a: b for a, b in w["opt"][k]["deltas"].items() if b} != {a: b for a, b in m["optionen"][k]["deltas"].items() if b}: fehl.append("%s%s Deltas Word %s Daten %s" % (mid, k, w["opt"][k]["deltas"], m["optionen"][k]["deltas"]))
        gl = " / ".join(norm(x) for x in m["glueck"])
        if norm(w["glueck"]) != norm(gl) and not (w["glueck"] == "" and not m["glueck"]):
            fehl.append(mid + " Glücksfaktor: Word %r Daten %r" % (w["glueck"][:40], gl[:40]))
    E.pruefe("Word gegen Daten: Mikroszenen (Titel, Vorlesen, Optionen, Wertaenderungen A/B, Glücksfaktoren)", G, not fehl, "; ".join(fehl[:5]))
    # Excel-Konstanten
    z = R.zeilen_nach_episode()
    ws = openpyxl.load_workbook(R.X.QUELLE_XLSX)["Tabelle1"]
    fehl = []
    for (i, k), info in z.items():
        r = info["zeile"]; o = eps_d[i]["optionen"][k]
        vgl = [("soll", o["soll"], ws["P%d" % r].value), ("zeit", o["aenderung"]["zeit"], ws["AH%d" % r].value), ("soz", o["aenderung"]["soz"], ws["AJ%d" % r].value),
               ("erfuellung", o["aenderung"]["erfuellung"], ws["AK%d" % r].value), ("lebensstil", o["aenderung"]["lebensstil"], ws["AL%d" % r].value),
               ("finanzen", o["aenderung"]["finanzen"], ws["AM%d" % r].value), ("energieBasis", o["energieBasis"], ws["AW%d" % r].value),
               ("beginn zeit", eps_d[i]["beginn"]["zeit"], ws["C%d" % r].value), ("beginn energie", eps_d[i]["beginn"]["energie"], ws["D%d" % r].value), ("beginn stimmung", eps_d[i]["beginn"]["stimmung"], ws["E%d" % r].value)]
        fehl += ["%s/%s %s: Daten %s Excel %s" % (i, k, n, a, b) for n, a, b in vgl if abs(a - b) > 1e-12]
        if o["excelZeile"] != r: fehl.append("%s/%s Zeile %s statt %s" % (i, k, o["excelZeile"], r))
    E.pruefe("Excel gegen Daten: Soll-Wert, 5 Aenderungen, Energie-Basis, Beginn-Aenderungen in allen 38 Zeilen", G, not fehl, "; ".join(fehl[:4]))
    kf = d["konfig"]
    gew = {"zeit": ws["AP3"].value, "energie": ws["AQ3"].value, "soz": ws["AR3"].value, "erfuellung": ws["AS3"].value, "lebensstil": ws["AT3"].value, "finanzen": ws["AU3"].value}
    E.pruefe("Gewichte in der Konfiguration = Excel AP3:AU3 (Zeile 3, in allen Zeilen gleich)", G, all(abs(kf["gewichte"][k] - gew[k]) < 1e-9 for k in gew), str(kf["gewichte"]))
    E.pruefe("Krisenepisoden 12A1-12B2: Beginn Stimmung -10, Soll 16 bis 20", G, all(eps_d[i]["beginn"]["stimmung"] == -10 and all(16 <= o["soll"] <= 20 for o in eps_d[i]["optionen"].values()) for i in ("12A1", "12A2", "12B1", "12B2")), "")
    # Wege
    def wege(i, pfad):
        out = []
        for k in "AB":
            n = eps_d[i]["optionen"][k]["weiter"]
            out += [pfad + [i + k]] if n is None else wege(n, pfad + [i + k])
        return out
    w = wege("1", [])
    E.pruefe("4096 moegliche Wege, alle genau 12 Episoden lang, alle 19 Episoden erreichbar", G, len(w) == 4096 and all(len(x) == 12 for x in w) and {s[:-1] for x in w for s in x} == set(ids), "%d Wege" % len(w))
    # Generator reproduzierbar
    tmp = os.path.join(tempfile.mkdtemp(prefix="p2-gen-"), "spiel-daten.js")
    p = subprocess.run([sys.executable, "-I", os.path.join(PROJ, "tools", "erzeuge_daten.py")], env=dict(os.environ, P2_AUSGABE=tmp), capture_output=True, text=True)
    E.pruefe("tools/erzeuge_daten.py erzeugt spiel-daten.js byte-identisch neu", G, p.returncode == 0 and os.path.exists(tmp) and open(tmp, "rb").read() == open(os.path.join(PROJ, "spiel-daten.js"), "rb").read(), p.stderr[-150:])
    # Handgeschriebene Texte (spiel-texte.js) gemeinsam mit den erzeugten Daten pruefen
    texte_pruefen(d, eps_d, ids)
    # Dateien
    gross = sum(os.path.getsize(os.path.join(PROJ, f)) for f in ("index.html", "spiel-daten.js", "spiel-texte.js", "spiel-kern.js"))
    E.pruefe("Gesamtgroesse index.html + Daten + Texte + Kern unter 300 KB (ohne Bilder)", G, gross < 300 * 1024, "%.0f KB" % (gross / 1024))
    html = open(os.path.join(PROJ, "index.html"), encoding="utf-8").read()
    alle = html + open(os.path.join(PROJ, "spiel-daten.js"), encoding="utf-8").read() + open(os.path.join(PROJ, "spiel-kern.js"), encoding="utf-8").read() + open(os.path.join(PROJ, "spiel-texte.js"), encoding="utf-8").read()
    E.pruefe("Keine externen Adressen (http/https/CDN) im Code ausser Namespace", G, not re.findall(r"https?://(?!www\.w3\.org)", alle), str(re.findall(r"https?://\S+", alle)[:3]))
    E.pruefe("Klassische Skripte: kein type=module, kein import/export, kein fetch/XMLHttpRequest", G,
             'type="module"' not in html and not re.search(r"^\s*(import|export)\s", alle, re.M) and not re.search(r"\bfetch\(|XMLHttpRequest", alle), "")
    E.pruefe("Relative Pfade: shared/base.css, spiel-daten.js, spiel-texte.js, spiel-kern.js eingebunden", G, '../shared/base.css' in html and 'src="spiel-daten.js"' in html and 'src="spiel-texte.js"' in html and 'src="spiel-kern.js"' in html, "")
    E.pruefe("lang=\"de\", Viewport-Meta, kein Eszett in Quelltext und Daten", G, 'lang="de"' in html and "viewport" in html and chr(0xDF) not in alle, "")
    for verboten in ("Pädagogische Kernidee", "Mini-Debriefing", "Wissenschaftliche Basis"):
        E.pruefe("Verbotener Inhalt nicht im Quelltext: «%s»" % verboten, G, verboten not in alle, "")
    E.speichern(os.path.join(HIER, "ergebnis_dateien.md"))
    print("Fehler:", len(E.fehler))
    return 1 if E.fehler else 0


if __name__ == "__main__":
    sys.exit(main())
