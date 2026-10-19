"""Unabhaengige Referenz fuer ganze Spielablaeufe: liest Parameter, Zeilenzuordnung und Beginn-Aenderungen
direkt aus dem Excel (nicht aus spiel-daten.js) und rechnet mit exakter Dezimalarithmetik (excel_referenz.rechne_decimal).
Verzweigung (Weiter bei) steht als Tabelle hier, Quelle: Word, Abschnitt je Episode; test_dateien.py vergleicht sie mit der Seite."""
import re
from decimal import Decimal

import openpyxl

import excel_referenz as X

WEITER = {("1", "A"): "2A", ("1", "B"): "2B", ("2A", "A"): "3", ("2A", "B"): "3", ("2B", "A"): "3", ("2B", "B"): "3",
          ("3", "A"): "4", ("3", "B"): "4", ("4", "A"): "5A", ("4", "B"): "5B",
          ("5A", "A"): "6", ("5A", "B"): "6", ("5B", "A"): "6", ("5B", "B"): "6", ("6", "A"): "7", ("6", "B"): "7",
          ("7", "A"): "8A", ("7", "B"): "8B", ("8A", "A"): "9", ("8A", "B"): "9", ("8B", "A"): "9", ("8B", "B"): "9",
          ("9", "A"): "10", ("9", "B"): "10", ("10", "A"): "11A", ("10", "B"): "11B",
          ("11A", "A"): "12A1", ("11A", "B"): "12A2", ("11B", "A"): "12B1", ("11B", "B"): "12B2"}


# Einordnung der Regeloptionen (unabhaengige Tabelle fuer den Test; Quelle: eigene Analyse der Optionen und des Word,
# siehe README «Ausrichtung»). L = langfristige Investition, K = kurzfristige Optimierung. Ausnahme: 11B (A = L, B = K).
AUSRICHTUNG = {ep: {"A": "kurzfristig", "B": "langfristig"} for ep in
               ("1", "2A", "2B", "3", "4", "5A", "5B", "6", "7", "8A", "8B", "9", "10", "11A", "12A1", "12A2", "12B1", "12B2")}
AUSRICHTUNG["11B"] = {"A": "langfristig", "B": "kurzfristig"}
UNSCHARF = {"6", "8A", "8B", "11B", "12A2", "12B2"}  # Episoden, bei denen die Zuordnung nicht eindeutig ist


def anzeige_zu_regel(tausch, buchstabe):
    """Angezeigter Buchstabe -> Regeloption (Excel-Zeile). Bei Tausch sind A und B vertauscht."""
    return {"A": "B", "B": "A"}[buchstabe] if tausch else buchstabe


def regel_zu_anzeige(tausch, regel):
    return anzeige_zu_regel(tausch, regel)  # symmetrisch


def zeilen_nach_episode():
    ws = openpyxl.load_workbook(X.QUELLE_XLSX)["Tabelle1"]
    out = {}
    for r in X.ZEILEN:
        m = re.search(r"Ep(\w+?) .* ([AB]) – ", ws["A%d" % r].value)
        out[(m.group(1), m.group(2))] = {"zeile": r, "C": Decimal(ws["C%d" % r].value), "D": Decimal(ws["D%d" % r].value),
                                         "E": Decimal(ws["E%d" % r].value)}
    return out


class Referenz:
    def __init__(self):
        self.params = X.lies_zeilenparameter()
        self.zeilen = zeilen_nach_episode()
        self.w = {k: Decimal(v) for k, v in dict(zeit=0, energie=0, soz=1, erfuellung=1, lebensstil=1, finanzen=1, stimmung=50).items()}
        self.log = []

    def beginn(self, ep):
        z = self.zeilen[(ep, "A")]
        self.w["zeit"] += z["C"]; self.w["energie"] += z["D"]; self.w["stimmung"] += z["E"]
        return dict(self.w)

    def tie(self, ep, opt, wurf):
        return X.ist_haelftenfall(self.params[self.zeilen[(ep, opt)]["zeile"]], dict(self.w, wurf=wurf))

    def episode(self, ep, opt, wurf):
        z = self.zeilen[(ep, opt)]
        r = X.rechne_decimal(self.params[z["zeile"]], dict(self.w, wurf=wurf))
        self.w = {k: r["nachher_" + k] for k in ("zeit", "energie", "soz", "erfuellung", "lebensstil", "finanzen", "stimmung")}
        self.log.append((ep, opt, wurf, r))
        return r

    def mikro(self, deltas):
        for k, v in deltas.items():
            if k in self.w and isinstance(v, (int, float)):
                self.w[k] += Decimal(str(v))
        return dict(self.w)

    def zahlen(self):
        return {k: float(v) for k, v in self.w.items()}
