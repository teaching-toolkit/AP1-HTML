#!/usr/bin/env python3
"""Erzeugt spiel-daten.js aus Excel und Word (keine Handabschrift).

Aufruf (aus dem Projektordner):
    python3 -I tools/erzeuge_daten.py [pfad/gluecksfaktoren.xlsx] [pfad/gluecksfaktor_spiel.docx]

Quellen (nur lesen):
  Excel  Blatt «Tabelle1», Zeilen 3 bis 46 (38 Zeilen = 19 Episoden x Option A/B)
         Zahlen: Spalten C-E (Beginn), G-M (Startwerte Zeile 3/4), P (Soll), AH-AM und AW
         (Veraenderungen), AP-AU (Gewichte).
         Formeln: Q, S-Y, AA-AF, AI, AN. Sie werden nicht uebernommen, sondern gegen die
         Muster geprueft, die spiel-kern.js umsetzt. Weicht eine Formel ab, bricht das
         Skript ab (ausser den zwei bekannten fehlenden Y-Formeln, die gemeldet werden).
  Word   Episoden (Vorlesen, Optionen, Unmittelbare Folge, Wuerfeltitel, Weiter bei,
         Glücksfaktoren), Mikroszenen M1-M10 mit Deltas, Zuordnung der Mikroszenen.
"""
import json
import os
import re
import sys
from collections import Counter, OrderedDict

import docx
import openpyxl
from openpyxl.utils import get_column_letter as L

HIER = os.path.dirname(os.path.abspath(__file__))
PROJEKT = os.path.dirname(HIER)
QUELLE = "/home/claude/html-projekte/quellen/p2"
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(QUELLE, "gluecksfaktoren.xlsx")
DOCX = sys.argv[2] if len(sys.argv) > 2 else os.path.join(QUELLE, "gluecksfaktor_spiel.docx")
AUSGABE = os.environ.get("P2_AUSGABE") or os.path.join(PROJEKT, "spiel-daten.js")  # Umgebungsvariable nur fuer den Reproduzierbarkeitstest

ZEILEN = list(range(3, 11)) + list(range(13, 21)) + list(range(23, 31)) + list(range(33, 47))

# ------------------------------------------------------------------ Excel
FORMELMUSTER = OrderedDict([
    ("Q", '=IF(O{r}<>"",IF(ROUND(2*SUM(AA{r}:AF{r})+O{r}, 1)<0, 0, ROUND(2*SUM(AA{r}:AF{r})+O{r}, 1)), "")'),
    ("S", '=IF(O{r}<>"", IF(L{r}<-2, G{r}+AH{r}-1, G{r}+AH{r}), "")'),
    ("T", '=IF(O{r}<>"",H{r}+AI{r},"")'),
    ("U", '=IF(O{r}<>"", I{r}+AJ{r}, "")'),
    ("V", '=IF(O{r}<>"", J{r}+AK{r}, "")'),
    ("W", '=IF(O{r}<>"", K{r}+AL{r}, "")'),
    ("X", '=IF(O{r}<>"", L{r}+AM{r}, "")'),
    ("Y", '=IF(O{r}<>"", M{r}+AN{r}, "")'),
    ("AA", '=IF(G{r}>2, AP{r}*2,IF(G{r}>4,AP{r}*3, G{r}*AP{r}))'),
    ("AB", '=IF(H{r}>2, AQ{r}*2,IF(H{r}>4,AQ{r}*3, H{r}*AQ{r}))'),
    ("AC", '=AR{r}*I{r}'),
    ("AD", '=AS{r}*J{r}'),
    ("AE", '=AT{r}*K{r}'),
    ("AF", '=IF(L{r}>2, AU{r}*2,IF(L{r}>4,AU{r}*3, L{r}*AU{r}))'),
    ("AI", '=AW{r}+ ROUND(8*(Y{r}-50)/50,0)/2'),
    ("AN", '=ROUND(8*(Q{r}-P{r})/(10), 0)'),
])
BEKANNT_FEHLEND = {("Y", 3), ("Y", 40)}  # Stimmung nach der Entscheidung: Formel fehlt


def zelle(ws, spalte, zeile):
    return ws["%s%d" % (spalte, zeile)].value


def lies_excel():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Tabelle1"]
    abweichungen = []
    fehlend = []
    for r in ZEILEN:
        for sp, muster in FORMELMUSTER.items():
            f = zelle(ws, sp, r)
            if f is None:
                fehlend.append((sp, r))
                continue
            if f != muster.format(r=r):
                abweichungen.append((sp, r, f))
    if abweichungen:
        raise SystemExit("Formel weicht vom erwarteten Muster ab: %r" % abweichungen[:5])
    if set(fehlend) != BEKANNT_FEHLEND:
        raise SystemExit("Unerwartet fehlende Formeln: %r" % sorted(set(fehlend) ^ BEKANNT_FEHLEND))

    # Gewichte: in allen Zeilen gleich (bis auf Gleitkommarauschen)
    gewicht_spalten = OrderedDict([("zeit", "AP"), ("energie", "AQ"), ("soz", "AR"),
                                   ("erfuellung", "AS"), ("lebensstil", "AT"), ("finanzen", "AU")])
    gewichte = OrderedDict()
    roh_zaehler = {}
    for key, sp in gewicht_spalten.items():
        werte = [zelle(ws, sp, r) for r in ZEILEN]
        ref = round(sum(werte) / len(werte), 10)
        if max(abs(w - ref) for w in werte) > 1e-12:
            raise SystemExit("Gewichte in Spalte %s sind nicht in allen Zeilen gleich" % sp)
        gewichte[key] = ref
        roh_zaehler[sp] = Counter(werte)

    start = OrderedDict()
    for key, sp in [("zeit", "G"), ("energie", "H"), ("soz", "I"), ("erfuellung", "J"),
                    ("lebensstil", "K"), ("finanzen", "L"), ("stimmung", "M")]:
        a, b = zelle(ws, sp, 3), zelle(ws, sp, 4)
        if a != b:
            raise SystemExit("Startwerte Zeile 3 und 4 unterschiedlich in Spalte %s" % sp)
        start[key] = a
    # Startwerte duerfen nur in Zeile 3/4 stehen
    for r in ZEILEN[2:]:
        for sp in "GHIJKLM":
            if zelle(ws, sp, r) is not None:
                raise SystemExit("Unerwarteter Wert in %s%d" % (sp, r))

    zeilen = OrderedDict()
    for r in ZEILEN:
        name = zelle(ws, "A", r)
        m = re.match(r"^Block (\d) – Ep(\d+[AB]?\d?) (.+?)\s+([AB]) – (.+)$", name)
        if not m:
            raise SystemExit("Unerwarteter Zeilenname: %r" % name)
        block, ep, ep_titel, opt, opt_name = m.groups()
        zeilen[r] = OrderedDict([
            ("zeile", r), ("name", name), ("block", int(block)), ("ep", ep), ("epTitel", ep_titel.strip()),
            ("option", opt), ("optName", opt_name.strip()),
            ("beginn", OrderedDict([("zeit", zelle(ws, "C", r)), ("energie", zelle(ws, "D", r)),
                                    ("stimmung", zelle(ws, "E", r))])),
            ("soll", zelle(ws, "P", r)),
            ("aenderung", OrderedDict([("zeit", zelle(ws, "AH", r)), ("soz", zelle(ws, "AJ", r)),
                                       ("erfuellung", zelle(ws, "AK", r)), ("lebensstil", zelle(ws, "AL", r)),
                                       ("finanzen", zelle(ws, "AM", r))])),
            ("energieBasis", zelle(ws, "AW", r)),
        ])
        for k, v in list(zeilen[r]["beginn"].items()) + list(zeilen[r]["aenderung"].items()) + [("soll", zeilen[r]["soll"]), ("eb", zeilen[r]["energieBasis"])]:
            if not isinstance(v, (int, float)):
                raise SystemExit("Keine Zahl in Zeile %d (%s): %r" % (r, k, v))
    return start, gewichte, zeilen, roh_zaehler


# ------------------------------------------------------------------ Word
def normalisiere(t):
    t = t.replace(" ", " ")
    t = re.sub(r"[ \t]+", " ", t).strip()
    return t


def anfuehrung_reparieren(t):
    """Gerade Schlusszeichen nach einem deutschen Anfuehrungszeichen durch die
    typografisch passenden ersetzen. Kein Wort wird geaendert."""
    t = re.sub(r"(„[^„“”\"]*)\"", lambda m: m.group(1) + "“", t)
    t = re.sub(r"(„[^„“”\"]*)”", lambda m: m.group(1) + "“", t)
    # Paare gerader Zeichen ("...") ohne oeffnendes „ ebenfalls typografisch setzen
    t = re.sub(r"\"([^\"„“”]*)\"", lambda m: "„" + m.group(1) + "“", t)
    return t


def aussen_anfuehrung_entfernen(t):
    t = t.strip()
    if t.startswith("„"):
        t = t[1:]
    t = re.sub(r"[“”\"]+$", "", t)
    return t.strip()


def zahl(s):
    s = s.replace("−", "-").replace("–", "-").replace(",", ".")
    v = float(s)
    return int(v) if v == int(v) else v


MIKRO_NAMEN = OrderedDict([
    ("Zeit", "zeit"), ("Energie", "energie"), ("Soziale Verbindungen", "soz"),
    ("Persönliche Erfüllung", "erfuellung"), ("Lebensstil-Qualität", "lebensstil"),
    ("Einkommen und Vermögen", "finanzen"),
])


def parse_deltas(text):
    text = text.strip().rstrip(".").strip()
    teile = [t.strip() for t in text.split(",") if t.strip()]
    out = OrderedDict()
    for t in teile:
        m = re.match(r"^(Zeit|Energie|Soziale Verbindungen|Persönliche Erfüllung|Lebensstil-Qualität|Einkommen und Vermögen)\s+([+\-−–]?\d+(?:[.,]\d+)?)$", t)
        if not m:
            raise SystemExit("Delta nicht lesbar: %r" % t)
        out[MIKRO_NAMEN[m.group(1)]] = zahl(m.group(2))
    return out


def lies_word():
    d = docx.Document(DOCX)
    absaetze = [(p.style.name, normalisiere(p.text)) for p in d.paragraphs]

    # ---- Episoden
    episoden = OrderedDict()
    block = None
    block_namen = {}
    ep = None
    opt = None
    modus = None
    i = 0
    start_idx = next(k for k, (s, t) in enumerate(absaetze) if t.startswith("BLOCK 1"))
    ende_idx = next(k for k, (s, t) in enumerate(absaetze) if s == "Heading 1" and t == "DEBRIEFING")
    for k in range(start_idx, ende_idx):
        stil, t = absaetze[k]
        if not t:
            modus = None if modus == "glueck" else modus
            continue
        if stil == "Heading 1":
            m = re.match(r"^BLOCK (\d) – (.+)$", t)
            block = int(m.group(1))
            block_namen[block] = m.group(2).strip()
            continue
        if stil.startswith("Heading 2"):
            m = re.match(r"^Episode\s+(\d+[AB]?\d?)\s*[:–-]\s*(.+)$", t)
            if m:
                ep = OrderedDict([("id", m.group(1)), ("block", block), ("titel", m.group(2).strip()),
                                  ("vorlesen", None), ("optionen", OrderedDict()), ("glueck", []), ("erkenntnis", []),
                                  ("hinweis", None)])
                episoden[ep["id"]] = ep
                opt = None
                modus = None
            elif t.startswith("(nur nach"):
                ep["hinweis"] = t
            else:
                raise SystemExit("Unbekannte Ueberschrift: %r" % t)
            continue
        if ep is None:
            continue
        if modus == "vorlesen_naechste":
            ep["vorlesen"] = t
            modus = None
            continue
        if t == "Vorlesen:":
            modus = "vorlesen_naechste"
            continue
        m = re.match(r"^Vorlesen:\s*(.+)$", t)
        if m:
            ep["vorlesen"] = m.group(1)
            continue
        m = re.match(r"^([AB])\s+(.+)$", t)
        if m and modus != "glueck":
            opt = OrderedDict([("label", m.group(2).strip()), ("folge", None), ("wuerfeltitel", None), ("weiter", None)])
            ep["optionen"][m.group(1)] = opt
            continue
        m = re.match(r"^Unmittelbare Folge:\s*(.+)$", t)
        if m:
            opt["folge"] = m.group(1)
            continue
        m = re.match(r"^Würfeln\s+(.+)$", t)
        if m and modus != "glueck":
            opt["wuerfeltitel"] = m.group(1).strip()
            continue
        m = re.match(r"^Weiter(?: bei)?:\s*(.+)$", t)
        if m:
            ziel = m.group(1)
            if "Ende" in ziel:
                opt["weiter"] = None
            else:
                mm = re.search(r"Episode\s+(\d+[AB]?\d?)", ziel)
                if not mm:
                    raise SystemExit("Weiter-Ziel nicht lesbar: %r" % ziel)
                opt["weiter"] = mm.group(1)
            continue
        if t == "Glücksfaktoren:":
            modus = "glueck"
            continue
        if modus == "glueck":
            if t.startswith("Wichtige Erkenntnis:") or t.startswith("Forschungserkenntnis:"):
                ep["erkenntnis"].append(t)
            else:
                ep["glueck"].append(t)
            continue
        raise SystemExit("Absatz nicht zugeordnet (Episode %s): %r" % (ep["id"], t))

    # ---- Mikroszenen
    mikro = OrderedDict()
    s0 = next(k for k, (s, t) in enumerate(absaetze) if t == "Mögliche Mikro-Szenen")
    s1 = next(k for k, (s, t) in enumerate(absaetze) if t.startswith("Das Prinzip ist immer"))
    cur = None
    for k in range(s0 + 1, s1):
        stil, t = absaetze[k]
        if not t:
            continue
        m = re.match(r"^M(\d+):\s*(.+)$", t)
        if m:
            cur = OrderedDict([("id", "M" + m.group(1)), ("titel", m.group(2).strip()), ("vorlesen", None),
                               ("optionen", OrderedDict()), ("glueck", [])])
            mikro[cur["id"]] = cur
            continue
        m = re.match(r"^Vorlesen:\s*(.+)$", t)
        if m:
            cur["vorlesen"] = m.group(1)
            continue
        m = re.match(r"^([AB])\s+(.+)$", t)
        if m:
            teile = [x.strip() for x in m.group(2).split("→")]
            if len(teile) != 3 or not teile[2].startswith("Weiter zur nächsten geplanten Szene"):
                raise SystemExit("Mikro-Option nicht lesbar: %r" % t)
            cur["optionen"][m.group(1)] = OrderedDict([("label", teile[0]), ("deltas", parse_deltas(teile[1]))])
            continue
        m = re.match(r"^Glücksfaktor(?:en)?:\s*(.*)$", t)
        if m:
            rest = m.group(1).strip()
            cur["glueck"] = [x.strip() for x in re.split(r"\s*/\s*(?=[A-G]\d+\s*[:–-])", rest)] if rest else []
            continue
        raise SystemExit("Mikroszenen-Absatz nicht zugeordnet: %r" % t)

    # ---- Zuordnung «Mikroszenen strategisch einsetzen»
    z0 = next(k for k, (s, t) in enumerate(absaetze) if t == "Mikroszenen strategisch einsetzen")
    z1 = next(k for k, (s, t) in enumerate(absaetze) if t.startswith("Einleitung einer Mikroszene"))
    zuordnung = []
    for k in range(z0 + 1, z1):
        stil, t = absaetze[k]
        if not t or t.startswith("Beobachtet während"):
            continue
        m = re.match(r"^M(\d+)\s*\((.+?)\)(?:\s*[–-]\s*(.+))?$", t)
        if m:
            zuordnung[-1]["szenen"].append(OrderedDict([("id", "M" + m.group(1)), ("hinweis", (m.group(3) or "").strip())]))
            continue
        if t.endswith(":"):
            zuordnung.append(OrderedDict([("kriterium", t[:-1]), ("szenen", [])]))
            continue
        raise SystemExit("Zuordnung nicht lesbar: %r" % t)
    return episoden, block_namen, mikro, zuordnung


# ------------------------------------------------------------------ Zusammenfuehren
def main():
    start, gewichte, zeilen, roh = lies_excel()
    episoden, block_namen, mikro, zuordnung = lies_word()

    # Episoden: Reihenfolge und Codes muessen mit dem Excel uebereinstimmen
    excel_eps = []
    for z in zeilen.values():
        if z["ep"] not in excel_eps:
            excel_eps.append(z["ep"])
    if excel_eps != list(episoden.keys()):
        raise SystemExit("Episodenreihenfolge Excel %r != Word %r" % (excel_eps, list(episoden.keys())))
    if len(episoden) != 19 or len(zeilen) != 38:
        raise SystemExit("Erwartet 19 Episoden / 38 Zeilen")
    if len(mikro) != 10:
        raise SystemExit("Erwartet 10 Mikroszenen, gefunden %d" % len(mikro))

    ausgabe_eps = []
    for ep_id, ep in episoden.items():
        zA = next(z for z in zeilen.values() if z["ep"] == ep_id and z["option"] == "A")
        zB = next(z for z in zeilen.values() if z["ep"] == ep_id and z["option"] == "B")
        if zA["beginn"] != zB["beginn"]:
            raise SystemExit("Beginn-Aenderungen unterscheiden sich in Episode %s: %r %r" % (ep_id, zA["beginn"], zB["beginn"]))
        if zA["block"] != ep["block"]:
            raise SystemExit("Block weicht ab in Episode %s" % ep_id)
        if sorted(ep["optionen"].keys()) != ["A", "B"]:
            raise SystemExit("Episode %s hat nicht genau die Optionen A und B" % ep_id)
        o = OrderedDict()
        for key, z in (("A", zA), ("B", zB)):
            w = ep["optionen"][key]
            for feld in ("folge", "wuerfeltitel"):
                if not w[feld]:
                    raise SystemExit("Episode %s Option %s: %s fehlt" % (ep_id, key, feld))
            if "weiter" not in w:
                raise SystemExit("Episode %s Option %s: Weiter fehlt" % (ep_id, key))
            o[key] = OrderedDict([
                ("label", anfuehrung_reparieren(w["label"])),
                ("folge", anfuehrung_reparieren(w["folge"])),
                ("wuerfeltitel", anfuehrung_reparieren(w["wuerfeltitel"])),
                ("weiter", w["weiter"]),
                ("soll", z["soll"]),
                ("aenderung", z["aenderung"]),
                ("energieBasis", z["energieBasis"]),
                ("excelZeile", z["zeile"]),
                ("excelName", z["name"]),
            ])
        ausgabe_eps.append(OrderedDict([
            ("id", ep_id), ("block", ep["block"]), ("titel", anfuehrung_reparieren(ep["titel"])),
            ("vorlesen", anfuehrung_reparieren(aussen_anfuehrung_entfernen(ep["vorlesen"]))),
            ("beginn", zA["beginn"]),
            ("optionen", o),
            ("glueck", [anfuehrung_reparieren(x) for x in ep["glueck"]]),
            ("erkenntnis", [anfuehrung_reparieren(x) for x in ep["erkenntnis"]]),
        ]))
        if not ep["vorlesen"]:
            raise SystemExit("Episode %s ohne Vorlesen-Text" % ep_id)

    ausgabe_mikro = []
    for m in mikro.values():
        if sorted(m["optionen"].keys()) != ["A", "B"] or not m["vorlesen"]:
            raise SystemExit("Mikroszene %s unvollstaendig" % m["id"])
        ausgabe_mikro.append(OrderedDict([
            ("id", m["id"]), ("titel", anfuehrung_reparieren(m["titel"])),
            ("vorlesen", anfuehrung_reparieren(aussen_anfuehrung_entfernen(m["vorlesen"]))),
            ("optionen", OrderedDict((k, OrderedDict([("label", anfuehrung_reparieren(v["label"])), ("deltas", v["deltas"])]))
                                     for k, v in m["optionen"].items())),
            ("glueck", [anfuehrung_reparieren(x) for x in m["glueck"]]),
        ]))

    # Weiter-Ziele muessen existieren
    for ep in ausgabe_eps:
        for k, o in ep["optionen"].items():
            if o["weiter"] is not None and o["weiter"] not in episoden:
                raise SystemExit("Unbekanntes Weiter-Ziel %r in %s%s" % (o["weiter"], ep["id"], k))

    attribute = [
        OrderedDict([("key", "zeit"), ("name", "Zeit"), ("langname", "Zeit")]),
        OrderedDict([("key", "energie"), ("name", "Energie"), ("langname", "Energie")]),
        OrderedDict([("key", "soz"), ("name", "Soz. Verb."), ("langname", "Soziale Verbindungen")]),
        OrderedDict([("key", "erfuellung"), ("name", "Erfüllung"), ("langname", "Persönliche Erfüllung")]),
        OrderedDict([("key", "lebensstil"), ("name", "Lebensstil"), ("langname", "Lebensstil-Qualität")]),
        OrderedDict([("key", "finanzen"), ("name", "Finanzen"), ("langname", "Finanzen")]),
    ]

    konfig = OrderedDict([
        ("_hinweis", "Zahlen und Formelparameter stammen aus gluecksfaktoren.xlsx (Tabelle1). Aenderungen hier wirken direkt auf das Spiel."),
        ("start", start),
        ("gewichte", gewichte),
        # Excel AA, AB, AF:  =IF(G>2, w*2, IF(G>4, w*3, G*w))  -- in dieser Reihenfolge ausgewertet.
        # Die zweite Stufe (ueber 4, Faktor 3) ist wegen der Reihenfolge nie erreichbar:
        # jeder Wert ueber 4 ist schon ueber 2 und wird in der ersten Stufe mit 2 abgefangen.
        # Folge: der positive Beitrag sattelt bei 2*Gewicht, negative Werte nicht.
        # Soz, Erfuellung, Lebensstil (AC, AD, AE) haben keine Saettigung.
        ("saettigungExcel", OrderedDict([
            ("zeit", [OrderedDict([("ueber", 2), ("faktor", 2)]), OrderedDict([("ueber", 4), ("faktor", 3)])]),
            ("energie", [OrderedDict([("ueber", 2), ("faktor", 2)]), OrderedDict([("ueber", 4), ("faktor", 3)])]),
            ("soz", None), ("erfuellung", None), ("lebensstil", None),
            ("finanzen", [OrderedDict([("ueber", 2), ("faktor", 2)]), OrderedDict([("ueber", 4), ("faktor", 3)])]),
        ])),
        # Excel Q: =IF(ROUND(2*SUM(AA:AF)+O,1)<0, 0, ROUND(2*SUM(AA:AF)+O,1))
        ("resultat", OrderedDict([("faktorSumme", 2), ("stellen", 1), ("minimum", 0)])),
        # Excel AN: =ROUND(8*(Q-P)/(10),0)
        ("stimmungsaenderung", OrderedDict([("faktor", 8), ("teiler", 10), ("stellen", 0)])),
        # Excel AI: =AW+ROUND(8*(Y-50)/50,0)/2
        ("energieAusStimmung", OrderedDict([("faktor", 8), ("nullpunkt", 50), ("teiler", 50), ("stellen", 0), ("durch", 2)])),
        # Excel S: =IF(L<-2, G+AH-1, G+AH)   (L = Finanzen VOR der Entscheidung)
        ("zeitAbzugFinanzen", OrderedDict([("finanzenUnter", -2), ("abzug", 1)])),
        # Word «Mechanik-Uebersicht»: Attribute -10..+10, Stimmung 0..100, unter 0 Spielende.
        # Das Excel begrenzt nicht. Standard: aus (= Excel). Umschalten per ?schranken=1.
        ("schranken", OrderedDict([("aktiv", False), ("attributMin", -10), ("attributMax", 10), ("stimmungMax", 100)])),
        ("spielendeStimmungUnter", 0),
        ("wuerfelSeiten", 20),
        # Word «Mikroszenen strategisch einsetzen»: Zeit/Energie unter -5.
        # Fuer «vernachlaessigt/niedrig» nennt das Word keine Zahl: Annahme unter 0 (siehe README).
        ("mikroVorschlag", OrderedDict([("zeitEnergieUnter", -5), ("attributUnter", 0)])),
    ])

    daten = OrderedDict([
        ("meta", OrderedDict([
            ("titel", "Glücksfaktoren-Spiel"),
            ("figur", "Alexandra"),
            ("quelleExcel", "gluecksfaktoren.xlsx, Blatt Tabelle1, Zeilen 3-46"),
            ("quelleWord", "gluecksfaktor_spiel.docx"),
            ("erzeugtMit", "tools/erzeuge_daten.py"),
            ("excelAuffaelligkeiten", [
                "Zellen Y3 und Y40 (Stimmung nach der Entscheidung) enthalten keine Formel. Die Seite rechnet dort wie in allen anderen 36 Zeilen M+AN.",
            ]),
        ])),
        ("attribute", attribute),
        ("konfig", konfig),
        ("blockNamen", OrderedDict((str(k), v) for k, v in block_namen.items())),
        ("episoden", ausgabe_eps),
        ("mikroszenen", ausgabe_mikro),
        ("mikroZuordnung", zuordnung),
    ])

    js = "/* AUTOMATISCH ERZEUGT von tools/erzeuge_daten.py aus Excel und Word. Nicht von Hand abschreiben.\n"
    js += "   Zum Aendern: Quelle anpassen und das Skript neu laufen lassen (oder hier gezielt editieren). */\n"
    js += "var SPIEL_DATEN = " + json.dumps(daten, ensure_ascii=False, indent=1) + ";\n"
    with open(AUSGABE, "w", encoding="utf-8") as f:
        f.write(js)

    print("spiel-daten.js geschrieben:", AUSGABE, "(%d Bytes)" % len(js.encode("utf-8")))
    print("Episoden:", [e["id"] for e in ausgabe_eps])
    print("Mikroszenen:", [m["id"] for m in ausgabe_mikro])
    print("Gewichte:", dict(gewichte))
    print("Startwerte:", dict(start))
    print("Fehlende Excel-Formeln (gemeldet):", sorted(BEKANNT_FEHLEND))
    for ep in ausgabe_eps:
        print("%-5s B%d  %-40s beginn=%s" % (ep["id"], ep["block"], ep["titel"][:40], dict(ep["beginn"])))
        for k, o in ep["optionen"].items():
            print("      %s  Zeile %2d  Soll %2d  EB %+d  %s  | Excel: %s | weiter=%s" % (
                k, o["excelZeile"], o["soll"], o["energieBasis"], dict(o["aenderung"]), o["excelName"].split(" – ", 1)[-1][:30], o["weiter"]))


if __name__ == "__main__":
    main()
